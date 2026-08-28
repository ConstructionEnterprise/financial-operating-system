from pathlib import Path
from openpyxl import load_workbook

PATH = Path("/home/ubuntu/ceff-deliverables/Cedarwood_Flats_SKU_Level_Underwriting_Model.xlsx")


def main():
    formulas = load_workbook(PATH, data_only=False)
    values = load_workbook(PATH, data_only=True)
    checks = [
        ("Inputs - SKU Register", "F25"),
        ("Inputs - SKU Register", "F26"),
        ("Rent Roll", "E10"),
        ("Rent Roll", "I10"),
        ("Rent Roll", "L10"),
        ("Rent Roll", "M10"),
        ("Rent Roll", "E15"),
        ("Revenue & Operations", "E15"),
    ]
    for sheet, coordinate in checks:
        print(f"{sheet}!{coordinate}")
        print(f"  formula: {formulas[sheet][coordinate].value}")
        print(f"  cached:  {values[sheet][coordinate].value}")


if __name__ == "__main__":
    main()
