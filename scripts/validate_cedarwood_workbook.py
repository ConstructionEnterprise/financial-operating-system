from pathlib import Path
from openpyxl import load_workbook

PATH = Path("/home/ubuntu/ceff-deliverables/Cedarwood_Flats_SKU_Level_Underwriting_Model.xlsx")
EXPECTED_SHEETS = ["Summary", "Inputs - SKU Register", "WBS Budget", "Rent Roll", "Revenue & Operations", "Monthly Schedule", "Sources & Uses", "Debt Schedule", "Returns"]


def main():
    if not PATH.exists():
        raise SystemExit(f"Missing workbook: {PATH}")
    formulas = load_workbook(PATH, data_only=False)
    values = load_workbook(PATH, data_only=True)
    if formulas.sheetnames != EXPECTED_SHEETS:
        raise SystemExit(f"Unexpected worksheet set: {formulas.sheetnames}")

    wbs = formulas["WBS Budget"]
    wbs_entries = [wbs.cell(row=row, column=3).value for row in range(10, 110)]
    if len(wbs_entries) != 100 or any(not str(value).startswith("WBS-") for value in wbs_entries):
        raise SystemExit("WBS does not contain exactly 100 SKU-level development entries")
    wbs_amounts = [wbs.cell(row=row, column=11).value for row in range(10, 110)]
    if sum(wbs_amounts) != 40_000_000:
        raise SystemExit(f"WBS does not total $40M: {sum(wbs_amounts):,}")
    if any(wbs.cell(row=row, column=11).comment is None for row in range(10, 110)):
        raise SystemExit("One or more hardcoded WBS budget amounts lack a source comment")

    inputs = formulas["Inputs - SKU Register"]
    sku_rows = [inputs.cell(row=row, column=3).value for row in range(10, inputs.max_row + 1)]
    required_skus = [value for value in sku_rows if isinstance(value, str) and (value.startswith(("PRJ-", "FIN-", "REV-", "OPX-", "SCH-", "CAP-", "EXIT-", "RET-")))]
    if len(required_skus) < 130:
        raise SystemExit(f"Expected at least 130 required underwriting SKUs; found {len(required_skus)}")
    uncommented_inputs = []
    for row in range(10, inputs.max_row + 1):
        for col in range(7, 11):
            cell = inputs.cell(row=row, column=col)
            if cell.value not in (None, "") and cell.font.color and cell.font.color.type == "rgb" and cell.font.color.rgb.endswith("0000FF") and cell.comment is None:
                uncommented_inputs.append(cell.coordinate)
    if uncommented_inputs:
        raise SystemExit("Hardcoded inputs missing source comments: " + ", ".join(uncommented_inputs[:10]))

    monthly = formulas["Monthly Schedule"]
    return_months = formulas["Returns"]
    if [monthly.cell(row=row, column=3).value for row in range(10, 46)] != list(range(1, 37)):
        raise SystemExit("Monthly Schedule is missing one or more months")
    if [return_months.cell(row=row, column=3).value for row in range(10, 46)] != list(range(1, 37)):
        raise SystemExit("Returns schedule is missing one or more months")

    formula_count = 0
    for sheet in formulas.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_count += 1
                    if "#REF!" in cell.value:
                        raise SystemExit(f"Broken formula reference in {sheet.title}!{cell.coordinate}: {cell.value}")
    if formula_count < 250:
        raise SystemExit(f"Insufficient linked formulas: {formula_count}")

    errors = []
    for sheet in values.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("#"):
                    errors.append(f"{sheet.title}!{cell.coordinate}={cell.value}")
    if errors:
        raise SystemExit("Cached spreadsheet errors detected: " + "; ".join(errors[:10]))

    cached_wbs_total = values["WBS Budget"]["K110"].value
    if cached_wbs_total != 40_000_000:
        raise SystemExit(f"Cached WBS total did not recalculate to $40M: {cached_wbs_total}")

    print("Workbook validation passed")
    print(f"Sheets: {len(formulas.sheetnames)}")
    print("WBS entries: 100")
    print(f"Required SKU entries: {len(required_skus)}")
    print("Hardcoded input-source comments: validated")
    print(f"Linked formulas: {formula_count}")
    print(f"Cached WBS total: {cached_wbs_total:,}")
    print("Cached spreadsheet errors: 0")


if __name__ == "__main__":
    main()
